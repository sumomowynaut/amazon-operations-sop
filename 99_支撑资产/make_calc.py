# -*- coding: utf-8 -*-
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '计算器')
os.makedirs(ROOT, exist_ok=True)

# ---- styles ----
F_TITLE = Font(name='微软雅黑', size=14, bold=True, color='FFFFFF')
F_HDR   = Font(name='微软雅黑', size=11, bold=True, color='FFFFFF')
F_LABEL = Font(name='微软雅黑', size=10)
F_IN    = Font(name='微软雅黑', size=10, bold=True, color='000000')
F_OUT   = Font(name='微软雅黑', size=10, bold=True, color='1F4E00')
F_NOTE  = Font(name='微软雅黑', size=8.5, italic=True, color='808080')

FILL_TITLE = PatternFill('solid', fgColor='2B4C7E')
FILL_HDR   = PatternFill('solid', fgColor='4472C4')
FILL_IN    = PatternFill('solid', fgColor='FFF2CC')  # yellow
FILL_OUT   = PatternFill('solid', fgColor='E2EFDA')  # green
FILL_NOTE  = PatternFill('solid', fgColor='F2F2F2')

thin = Side(style='thin', color='BBBBBB')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

def setw(ws, widths):
    for i,w in enumerate(widths,1):
        ws.column_dimensions[get_column_letter(i)].width = w

def title(ws, text, ncol=3):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    c = ws.cell(1,1,text); c.font=F_TITLE; c.fill=FILL_TITLE
    c.alignment=Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[1].height=26

def note(ws, row, text, ncol=3):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncol)
    c=ws.cell(row,1,text); c.font=F_NOTE; c.fill=FILL_NOTE
    c.alignment=Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[row].height=28

def section(ws, row, text, ncol=3):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncol)
    c=ws.cell(row,1,text); c.font=F_HDR; c.fill=FILL_HDR
    c.alignment=Alignment(horizontal='left', vertical='center')

def inp(ws, row, label, default, note_txt=None, fmt=None):
    ws.cell(row,1,label).font=F_LABEL
    b=ws.cell(row,2,default); b.font=F_IN; b.fill=FILL_IN; b.border=BORDER
    if fmt: b.number_format=fmt
    ws.cell(row,1).border=BORDER
    if note_txt:
        c=ws.cell(row,3,note_txt); c.font=F_NOTE; c.border=BORDER
    return b

def out(ws, row, label, formula, note_txt=None, fmt=None):
    ws.cell(row,1,label).font=F_LABEL
    b=ws.cell(row,2,formula); b.font=F_OUT; b.fill=FILL_OUT; b.border=BORDER
    if fmt: b.number_format=fmt
    ws.cell(row,1).border=BORDER
    if note_txt:
        c=ws.cell(row,3,note_txt); c.font=F_NOTE; c.border=BORDER
    return b

CUR='$#,##0.00'; PCT='0.0%'; NUM='0'; NUM2='0.0'; DATE='yyyy-mm-dd'

# ================= C01 利润计算器 =================
wb=Workbook(); ws=wb.active; ws.title='C01利润计算器'
setw(ws,[30,14,46])
title(ws,'C01 利润计算器（全费用·四层利润口径）')
note(ws,2,'黄色=手动输入；绿色=公式自动计算。费率为教学默认值，正式使用请登录 Seller Central 核实当前值（证据等级 A/B，需定期验证）。口径与 09-01 一致。')
section(ws,4,'一、输入区（黄色单元格，可修改）')
inp(ws,5,'售价（美元/件）',21.99,'目标售价',CUR)
inp(ws,6,'佣金率 %',15,'A级：按类目，需登录核实',NUM2)
inp(ws,7,'FBA 配送费（美元/件）',6.30,'A级：按尺寸/重量分档+附加费，需登录核实',CUR)
inp(ws,8,'产品成本 COGS（美元/件）',3.80,'含包装，是否含头程由团队口径定',CUR)
inp(ws,9,'头程物流（美元/件）',0.90,'海运/空运分摊',CUR)
inp(ws,10,'广告费（美元/件）',3.50,'广告预算 ÷ 预计销量',CUR)
inp(ws,11,'Coupon/促销费率 %',5,'券面值占售价比例',NUM2)
inp(ws,12,'退货率 %',6,'D级：按类目/历史经验',NUM2)
inp(ws,13,'退货单件损失（美元/件）',1.40,'退款+退货处理费−残值（D口径）',CUR)
inp(ws,14,'仓储费（美元/件）',0.45,'月度+长期仓储分摊',CUR)
inp(ws,15,'固定成本分摊（美元/件）',0.30,'工资/房租/工具分摊',CUR)
inp(ws,16,'收款手续费率 %',1,'Payoneer/PingPong 等',NUM2)
inp(ws,17,'汇损率 %',0.5,'跨币种结算汇率差',NUM2)
inp(ws,18,'认证费摊销（美元/件）',0.20,'FDA/CE/UL/CPC 等',CUR)
inp(ws,19,'软件订阅费摊销（美元/件）',0.15,'Helium10 等工具',CUR)
inp(ws,20,'移除/弃置费摊销（美元/件）',0.05,'清库存成本',CUR)
inp(ws,21,'样品/打样费摊销（美元/件）',0.10,'新品开发期',CUR)
section(ws,23,'二、输出区（自动计算，勿手改）')
out(ws,24,'佣金金额','=B5*B6/100','售价×佣金率',CUR)
out(ws,25,'净收入','=B5-B24-B7','售价−佣金−FBA费',CUR)
out(ws,26,'毛利','=B25-B8-B9','净收入−成本−头程',CUR)
out(ws,27,'毛利率','=IF(B5=0,0,B26/B5)','毛利/售价',PCT)
out(ws,28,'Coupon 金额','=B5*B11/100','售价×Coupon费率',CUR)
out(ws,29,'贡献利润','=B26-B10-B28-B13','毛利−广告−Coupon−退货',CUR)
out(ws,30,'贡献利润率','=IF(B5=0,0,B29/B5)','贡献利润/售价',PCT)
out(ws,31,'运营利润','=B29-B14-B15','贡献−仓储−固定分摊',CUR)
out(ws,32,'收款手续费金额','=B5*B16/100','售价×手续费率',CUR)
out(ws,33,'汇损金额','=B5*B17/100','售价×汇损率',CUR)
out(ws,34,'净利润','=B31-B32-B33-B18-B19-B20-B21','运营−手续费−汇损−认证−软件−移除−样品',CUR)
out(ws,35,'净利率','=IF(B5=0,0,B34/B5)','净利润/售价',PCT)
out(ws,36,'盈亏平衡 ACOS','=IF(B5=0,0,(B5-B24-B7-B8-B9-B13)/B5)','广告上限，非目标值',PCT)
out(ws,37,'盈亏判断','=IF(B34>=0,"盈利","亏损")','',None)
note(ws,39,'使用说明：1) 所有黄色输入项均需按实际填写；2) 费率以 Seller Central 当前值为准；3) 盈亏平衡 ACOS 用于广告判断（见 05-08/05-12），超过即广告边际亏损。')
wb.save(os.path.join(ROOT,'C01_利润计算器.xlsx'))

# ================= C03 库存补货计算器 =================
wb=Workbook(); ws=wb.active; ws.title='C03库存补货'
setw(ws,[30,14,46])
title(ws,'C03 库存补货计算器')
note(ws,2,'黄色=输入；绿色=自动计算。公式为团队经验（D级，非官方）。断货预警线 14 天/关注线 30 天为 D 级经验值。')
section(ws,4,'一、输入区')
inp(ws,5,'日均销量（件/天）',20,'建议取近 30 天均值',NUM2)
inp(ws,6,'生产时间（天）',15,'供应商交期',NUM)
inp(ws,7,'运输时间（天）',25,'海运/空运',NUM)
inp(ws,8,'入仓时间（天）',5,'FBA 签收至上架',NUM)
inp(ws,9,'安全库存（件）',200,'覆盖波动；建议=日均销量×安全天数',NUM)
inp(ws,10,'当前 FBA 可售（件）',600,'Fulfillable',NUM)
inp(ws,11,'在途库存（件）',0,'已发货未上架',NUM)
inp(ws,12,'目标覆盖天数（天）',60,'想覆盖多久的销量',NUM)
section(ws,14,'二、输出区')
out(ws,15,'Lead Time（总补货周期，天）','=B6+B7+B8','生产+运输+入仓',NUM)
out(ws,16,'补货点（件）','=B5*B15+B9','日均×LeadTime+安全库存',NUM)
out(ws,17,'当前库存覆盖天数','=IF(B5=0,0,B10/B5)','可售/日均销量',NUM2)
out(ws,18,'预计断货日期','=IF(B5=0,"",TODAY()+B10/B5)','今天+覆盖天数',DATE)
out(ws,19,'建议补货量（件）','=MAX(0,B12*B5-B10-B11)','目标覆盖×日均−可售−在途',NUM)
out(ws,20,'断货风险判断','=IF(B5=0,"",IF(B17<14,"高风险：立即空运/调拨",IF(B17<30,"关注：尽快安排补货","健康")))','D级阈值 14/30 天',None)
note(ws,22,'补货端到端跟踪字段：下单日期 | 生产完成 | 发货 | 到港 | FBA 签收 | 上架可售 | 当前可售。每步更新一次（对应 07-02）。')
wb.save(os.path.join(ROOT,'C03_库存补货计算器.xlsx'))

# ================= C04 广告 ACOS/TACOS 计算器 =================
wb=Workbook(); ws=wb.active; ws.title='C04广告ACOS'
setw(ws,[30,14,46])
title(ws,'C04 广告 ACOS / TACOS 计算器')
note(ws,2,'黄色=输入；绿色=自动计算。TACOS 为非官方指标（D级）；ACOS/ROAS 为官方指标（B级）。')
section(ws,4,'一、输入区')
inp(ws,5,'广告花费 Spend（美元）',1600,'周期内总广告费',CUR)
inp(ws,6,'广告归因销售 Ad Sales（美元）',4000,'Ads Console 归因销售',CUR)
inp(ws,7,'自然销售 Organic Sales（美元）',2000,'总销售−广告销售',CUR)
inp(ws,8,'售价（美元/件）',21.99,'用于算盈亏平衡',CUR)
inp(ws,9,'佣金率 %',15,'A级，需登录核实',NUM2)
inp(ws,10,'FBA 费（美元/件）',6.30,'A级，需登录核实',CUR)
inp(ws,11,'产品成本（美元/件）',3.80,'含包装',CUR)
inp(ws,12,'头程（美元/件）',0.90,'',CUR)
inp(ws,13,'退货损失（美元/件）',1.40,'',CUR)
section(ws,15,'二、输出区')
out(ws,16,'总销售 Total Sales','=B6+B7','广告+自然',CUR)
out(ws,17,'ACOS','=IF(B6=0,0,B5/B6)','广告花费/广告销售',PCT)
out(ws,18,'ROAS','=IF(B5=0,0,B6/B5)','广告销售/广告花费',NUM2)
out(ws,19,'TACOS','=IF(B16=0,0,B5/B16)','广告花费/总销售（D级非官方）',PCT)
out(ws,20,'广告销售占比','=IF(B16=0,0,B6/B16)','近似广告订单占比',PCT)
out(ws,21,'盈亏平衡 ACOS','=IF(B8=0,0,(B8-B8*B9/100-B10-B11-B12-B13)/B8)','广告停止线',PCT)
out(ws,22,'广告盈亏判断','=IF(B17>B21,"超盈亏平衡线：广告边际亏损","低于盈亏平衡线：可接受")','',None)
note(ws,24,'判断要点：新品期看 TACOS+自然排名趋势（05-12），不孤立看 ACOS；ACOS 超过盈亏平衡线且 TACOS 也高=纯亏损，进 05-11 诊断。')
wb.save(os.path.join(ROOT,'C04_广告ACOS_TACOS计算器.xlsx'))

print('OK 3 files:')
for f in ['C01_利润计算器.xlsx','C03_库存补货计算器.xlsx','C04_广告ACOS_TACOS计算器.xlsx']:
    p=os.path.join(ROOT,f); print(' ', f, os.path.getsize(p),'bytes')
